from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).with_name("unit-02-activities-memory-units.docx")


def set_rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    set_rtl(p)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", *, style=None, bold=False, size=11, space_after=5):
    p = doc.add_paragraph(style=style)
    set_rtl(p)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(15 if level == 1 else 12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_rtl(p)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        run.font.size = Pt(10.5)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        if widths:
            cell.width = Cm(widths[i])
        set_cell_shading(cell, "1F4E78")
        set_cell_text(cell, header, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if widths:
                cells[i].width = Cm(widths[i])
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[i], "EAF2F8")
            set_cell_text(cells[i], value)
    doc.add_paragraph()
    return table


def add_rule_line(doc):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("........................................................................................................")
    r.font.size = Pt(10)
    return p


def add_footer(section):
    p = section.footer.paragraphs[0]
    set_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("الإعلام الآلي - السنة الأولى ثانوي - الوحدة 02")
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(8.5)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    add_footer(section)

    title = doc.add_paragraph(style="Title")
    set_rtl(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("نشاطات حصة وحدات قياس الذاكرة")
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(21)
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = doc.add_paragraph()
    set_rtl(subtitle)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    subrun = subtitle.add_run("دليل الأستاذ وورقة عمل التلميذ - الوحدة 02 مكونات الحاسوب وتجميعه")
    subrun.font.name = "Arial"
    subrun._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    subrun.font.size = Pt(11)

    add_heading(doc, "بطاقة الحصة", 2)
    add_table(doc, ["العنصر", "التفصيل"], [
        ("المدة", "55 دقيقة"),
        ("المكتسبات السابقة", "وحدات الإدخال والإخراج، الوحدة المركزية، والفرق بين RAM والتخزين الدائم."),
        ("الأهداف", "ترتيب وحدات السعة، التحويل بين MB وGB، وربط السعة بموقف واقعي."),
        ("الوسائل", "سبورة، بطاقات ورقية، أوراق صغيرة، أو عرض صورة سلم وحدات السعة."),
    ], [4.3, 12.5])

    add_heading(doc, "سير الحصة", 2)
    add_table(doc, ["المرحلة", "الزمن", "تعليمة الأستاذ وناتج التلميذ"], [
        ("استرجاع من أنا", "8 د", "اعرض: لوحة مفاتيح، شاشة، معالج، RAM، قرص. يحدد التلاميذ التصنيف والوظيفة."),
        ("رتب سلم الذاكرة", "15 د", "تعمل الأفواج ببطاقات bit وByte وKB وMB وGB وTB، ثم تضع 8 و1024 في مواضعهما."),
        ("تحدي التحويلات", "15 د", "تحل الأفواج تحويلات قصيرة وتشرح طريقة واحدة أمام القسم."),
        ("ملفات النادي العلمي", "12 د", "تطبق الأفواج التحويل في وضعية تخزين صور وفيديو."),
        ("بطاقة خروج", "5 د", "يجيب كل تلميذ منفرداً عن ثلاثة أسئلة قصيرة."),
    ], [3.6, 2.0, 11.2])

    add_paragraph(doc, "تنبيه بيداغوجي: استعمل 1024 في جميع التحويلات داخل هذه الحصة، ولا تنتقل بعدُ إلى المنافذ أو النظام الثنائي.", bold=True, size=10)

    doc.add_page_break()
    title2 = doc.add_paragraph(style="Title")
    set_rtl(title2)
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title2.add_run("ورقة عمل التلميذ")
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(20)

    add_paragraph(doc, "الاسم واللقب: ...............................................    القسم: ...............    الفوج: ...............", size=10.5, space_after=10)
    add_heading(doc, "النشاط 1 رتب سلم الذاكرة", 2)
    add_paragraph(doc, "رتب الوحدات من الأصغر إلى الأكبر، ثم أكمل العلاقات:", size=11)
    add_rule_line(doc)
    add_rule_line(doc)
    add_paragraph(doc, "1 Byte = ........ bits     1 KB = ........ B     1 MB = ........ KB     1 GB = ........ MB", size=11, space_after=10)

    add_heading(doc, "النشاط 2 تحدي التحويلات", 2)
    add_table(doc, ["التحويل", "الحل والطريقة"], [
        ("2048 MB = ........ GB", ""),
        ("3 GB = ........ MB", ""),
        ("4096 KB = ........ MB", ""),
        ("2 TB = ........ GB", ""),
    ], [7.2, 9.6])

    add_heading(doc, "النشاط 3 ملفات النادي العلمي", 2)
    add_paragraph(doc, "يريد النادي العلمي حفظ 300 صورة، حجم كل صورة 4 MB، وفيديو حجمه 800 MB. هل تكفي ذاكرة USB سعتها 2 GB؟ علل إجابتك بالحساب.", size=11)
    for _ in range(4):
        add_rule_line(doc)

    add_heading(doc, "بطاقة الخروج", 2)
    add_paragraph(doc, "1. 1 GB = ........ MB      2. هل RAM تحفظ الملفات بعد إيقاف الحاسوب؟ لماذا؟      3. اذكر جهاز إدخال وجهاز إخراج.", size=10.5)
    add_rule_line(doc)
    add_rule_line(doc)

    doc.add_page_break()
    title3 = doc.add_paragraph(style="Title")
    set_rtl(title3)
    title3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title3.add_run("التصحيح ومعايير الملاحظة")
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(20)

    add_heading(doc, "التصحيح", 2)
    add_bullets(doc, [
        "سلم الوحدات: bit ثم Byte ثم KB ثم MB ثم GB ثم TB.",
        "1 Byte = 8 bits، وكل انتقال لاحق في السلم يساوي 1024 من الوحدة السابقة.",
        "2048 MB = 2 GB، و3 GB = 3072 MB، و4096 KB = 4 MB، و2 TB = 2048 GB.",
        "ملفات النادي: 300 × 4 MB = 1200 MB. الحجم الكلي 1200 + 800 = 2000 MB. و2 GB = 2048 MB، لذلك تكفي الذاكرة ويتبقى 48 MB تقريباً.",
        "بطاقة الخروج: 1 GB = 1024 MB. RAM ذاكرة مؤقتة. مثال: لوحة المفاتيح إدخال والشاشة إخراج.",
    ])

    add_heading(doc, "شبكة ملاحظة سريعة", 2)
    add_table(doc, ["المعيار", "متحقق", "بحاجة إلى دعم"], [
        ("يرتب الوحدات من الأصغر إلى الأكبر", "□", "□"),
        ("يستعمل 1024 في الاتجاه الصحيح", "□", "□"),
        ("ينجز تحويلاً بين MB وGB", "□", "□"),
        ("يشرح أن RAM ذاكرة مؤقتة", "□", "□"),
        ("يشارك في عمل الفوج ويحترم دوره", "□", "□"),
    ], [10.0, 3.4, 3.4])

    add_paragraph(doc, "إثراء اختياري للتلاميذ الأسرع: إذا كانت لديك 5 GB فارغة، فكم ملفاً حجمه 250 MB يمكن تخزينه؟ الإجابة: 20 ملفاً.", bold=True, size=10.5, space_after=0)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
